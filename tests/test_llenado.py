import datetime as dt
import io
import zipfile

import openpyxl

from karina import llenado


def _libro():
    """CONTABILIDAD minimo: JP y RC con dos fechas (fila 5 y 6), y una
    pestaña ajena que no se debe tocar."""
    wb = openpyxl.Workbook()
    wb.active.title = "SALDOS"
    wb["SALDOS"]["B2"] = "no tocar"
    for h in ("JP", "RC"):
        ws = wb.create_sheet(h)
        ws["A5"] = dt.datetime(2026, 10, 10)
        ws["A6"] = "=A5+2"
        for r in (5, 6):
            ws[f"C{r}"] = 100
            ws[f"I{r}"] = 10
            ws[f"X{r}"] = f"=C{r}+D{r}+E{r}+H{r}+F{r}+G{r}"
            ws[f"AL{r}"] = f"=AJ{r}-AK{r}"
        ws["O5"] = "=SUM(I5:L5,P5)"
        ws["O6"] = "=SUM(I6:L6,P6)+O5"   # formula vieja: deja fuera INDIA/MH
        ws["W5"] = "=SUM(Q5:V5)"
        ws["W6"] = "=SUM(Q6:V6)+W5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _dia(fecha, fila, nueva, esperado, difs=(), reinicio=False):
    h = {"fila": fila, "nueva": nueva, "reinicio": reinicio, "comite": False, "pregunta_alquiler": None,
         "esperado": esperado, "diferencias": list(difs)}
    return {"fecha": fecha, "avisos": [], "registros": 1, "hojas": {"JP": h, "RC": dict(h)}}


def _abrir(datos):
    return openpyxl.load_workbook(io.BytesIO(datos))


def test_fila_existente_solo_corrige_diferencias_y_formula_de_acumulado():
    datos = _libro()
    wf = _abrir(datos)
    dia = _dia(dt.date(2026, 10, 12), 6, False, {"C": 150.0, "I": 10.0, "O": 20.0, "W": 0.0},
               difs=[{"columna": "C", "esperado": 150.0}])
    salida, cambios = llenado.llenar(datos, [dia], wf, wf)
    ws = _abrir(salida)["JP"]
    assert ws["C6"].value == 150
    assert ws["I6"].value == 10
    assert ws["O6"].value == "=SUM(I6:N6,P6)+O5"
    assert ws["W6"].value == "=SUM(Q6:V6)+W5"  # ya era la correcta: no se toca
    assert cambios[0]["celdas"] == ["C", "O"]


def test_fila_nueva_se_crea_con_fecha_valores_y_formulas_corridas():
    datos = _libro()
    wf = _abrir(datos)
    wv = openpyxl.load_workbook(io.BytesIO(datos), data_only=True)
    # data_only de un libro guardado por openpyxl no tiene cache: la fecha
    # de la fila anterior se toma del dia anterior del rango.
    d12 = _dia(dt.date(2026, 10, 12), 6, False, {"C": 100.0}, difs=[])
    d13 = _dia(dt.date(2026, 10, 13), 7, True, {"C": 4645.0, "M": 35.0, "Z": 69280.0, "O": 0.0, "W": 0.0})
    salida, cambios = llenado.llenar(datos, [d12, d13], wf, wv)
    ws = _abrir(salida)["JP"]
    assert ws["C7"].value == 4645
    assert ws["M7"].value == 35
    assert ws["Z7"].value == 69280
    assert ws["O7"].value == "=SUM(I7:N7,P7)+O6"
    assert ws["W7"].value == "=SUM(Q7:V7)+W6"
    assert ws["X7"].value == "=C7+D7+E7+H7+F7+G7"
    assert ws["AL7"].value == "=AJ7-AK7"
    assert ws["A7"].value.startswith("=A6+")
    assert any(c["nueva"] for c in cambios)


def test_dia_de_reinicio_usa_formula_sin_acumulado_anterior():
    datos = _libro()
    wf = _abrir(datos)
    dia = _dia(dt.date(2026, 10, 12), 6, False, {"O": 0.0, "W": 0.0}, reinicio=True)
    ws = _abrir(llenado.llenar(datos, [dia], wf, wf)[0])["JP"]
    assert ws["O6"].value == "=SUM(I6:N6,P6)"


def test_no_toca_otras_pestanas_y_marca_recalculo():
    datos = _libro()
    wf = _abrir(datos)
    dia = _dia(dt.date(2026, 10, 12), 6, False, {"C": 1.0}, difs=[{"columna": "C", "esperado": 1.0}])
    salida, _ = llenado.llenar(datos, [dia], wf, wf)
    z_in, z_out = zipfile.ZipFile(io.BytesIO(datos)), zipfile.ZipFile(io.BytesIO(salida))
    rutas = llenado._rutas_hojas(z_in)
    assert z_in.read(rutas["SALDOS"]) == z_out.read(rutas["SALDOS"])
    assert b'fullCalcOnLoad="1"' in z_out.read("xl/workbook.xml")
    assert _abrir(salida)["SALDOS"]["B2"].value == "no tocar"


def test_no_pisa_la_celda_maestra_de_una_formula_compartida():
    from lxml import etree
    raiz = etree.fromstring(
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        b'<row r="6"><c r="X6"><f t="shared" ref="X6:X9" si="0">C6+D6</f><v>1</v></c></row>'
        b'</sheetData></worksheet>')
    hoja = llenado._Hoja(raiz)
    assert llenado._escribir(hoja, 6, "X", "f", "C6+D6+E6") is False
    assert llenado._escribir(hoja, 6, "C", "v", 12.5) is True
    assert hoja.celda(6, "C").find(llenado._q("v")).text == "12.5"
