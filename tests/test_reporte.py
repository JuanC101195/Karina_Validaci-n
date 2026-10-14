import datetime as dt
import io

import openpyxl

from karina import reporte


def _dia(diferencias, fila=479, avisos=()):
    hoja = {"fila": fila, "comite": False, "esperado": {"C": 4645.0, "AJ": 6764386.47, "Y": 0.0},
            "diferencias": diferencias}
    return {"fecha": dt.date(2026, 10, 10), "registros": 42, "avisos": list(avisos),
            "hojas": {"JP": hoja, "RC": dict(hoja, diferencias=[])}}


def _abrir(datos):
    return openpyxl.load_workbook(io.BytesIO(datos))


def test_genera_pestanas_jp_y_rc_con_columnas_de_contabilidad():
    wb = _abrir(reporte.generar_excel([_dia([])]))
    assert wb.sheetnames == ["JP", "RC"]
    ws = wb["JP"]
    assert ws["A3"].value.date() == dt.date(2026, 10, 10)
    assert ws["C3"].value == 4645.0
    assert ws["AJ3"].value == 6764386.47
    assert ws["X3"].value is None  # columna de formula en CONTABILIDAD: no se llena


def test_celda_con_diferencia_va_en_rojo_con_nota():
    dif = {"columna": "AJ", "concepto": "SAL CAJA 3", "excel": 6765387.47, "esperado": 6764386.47, "diferencia": 1001.0}
    ws = _abrir(reporte.generar_excel([_dia([dif])]))["JP"]
    assert ws["AJ3"].fill.fgColor.rgb.endswith("F8CBAD")
    assert "6,765,387.47" in ws["AJ3"].comment.text
    assert ws["C3"].fill.fgColor.rgb.endswith("E2EFDA")
    assert ws['AJ3'].comment is not None
    assert _abrir(reporte.generar_excel([_dia([dif])]))['RC']['AJ3'].comment is None


def test_avisos_de_una_hoja_no_aparecen_en_la_otra():
    wb = _abrir(reporte.generar_excel([_dia([], avisos=["JP: Pronto distinto"])]))
    assert "Pronto" in wb["JP"]["A3"].comment.text
    assert wb["RC"]["A3"].comment is None


def test_fecha_sin_fila_en_contabilidad():
    d = _dia([])
    d["hojas"]["RC"] = {"fila": None, "comite": False, "esperado": {}, "diferencias": []}
    ws = _abrir(reporte.generar_excel([d]))["RC"]
    assert "no tiene fila" in ws["A3"].comment.text


def test_alquiler_de_comite_pendiente_va_en_amarillo():
    d = _dia([])
    d["hojas"]["JP"]["pregunta_alquiler"] = {"total": 48840.0, "respuesta": None, "pendiente": True}
    ws = _abrir(reporte.generar_excel([d]))["JP"]
    assert ws['AB3'].fill.fgColor.rgb.endswith('FFF2CC')
    assert '48,840.00' in ws['AC3'].comment.text
