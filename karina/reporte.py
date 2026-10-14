"""
Excel de salida: pestañas JP y RC con el mismo orden de columnas que
CONTABILIDAD (A = fecha ... AJ = SAL CAJA 3), una fila por dia, con los
valores que salen de OpenCorp y Caja 3. Las celdas que no coinciden con lo
digitado van en rojo con una nota del valor que tenia CONTABILIDAD, para
que Karina copie bloques sin buscar a mano.

Las columnas que en CONTABILIDAD son formula (acumulados, totales,
chequeos) se dejan vacias: copiarlas encima borraria la formula.
"""
import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter

from .reglas import HOJAS, NOMBRES_COLUMNA

ULTIMA_COL = "AJ"
N_COLS = column_index_from_string(ULTIMA_COL)
COLUMNAS_FORMULA = ("B", "X", "AF", "AH", "AI")  # no se calculan aca
# O y W SI se calculan (para validar el acumulado), pero en CONTABILIDAD
# son formula: van en cursiva y el encabezado avisa que no se peguen.
COLUMNAS_SOLO_VALIDAR = ("O", "W")
AUTOR = "Validaciones"

ROJO = PatternFill("solid", fgColor="F8CBAD")
VERDE = PatternFill("solid", fgColor="E2EFDA")
GRIS = PatternFill("solid", fgColor="EDEDED")
ENCABEZADO = PatternFill("solid", fgColor="1F3864")
AMARILLO = PatternFill("solid", fgColor="FFF2CC")
AZUL = PatternFill("solid", fgColor="DDEBF7")  # fila nueva: el programa la llena
BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
FORMATO_NUM = "#,##0.0;-#,##0.0;0.0"


def _encabezados(encabezados, hoja, col):
    """Dos filas de encabezado tomadas del CONTABILIDAD subido (filas 2 y 3);
    si no vienen, el nombre del mapeo."""
    fila2, fila3 = (encabezados or {}).get(hoja, ({}, {}))
    arriba, abajo = fila2.get(col), fila3.get(col)
    if arriba is None and abajo is None:
        return None, NOMBRES_COLUMNA.get(col)
    return arriba, abajo


def _escribir_encabezados(ws, hoja, encabezados):
    for c in range(1, N_COLS + 1):
        col = get_column_letter(c)
        arriba, abajo = ("FECHA", None) if col == "A" else _encabezados(encabezados, hoja, col)
        for fila, texto in ((1, arriba), (2, abajo)):
            celda = ws.cell(fila, c, texto)
            celda.fill = ENCABEZADO
            celda.font = Font(bold=True, color="FFFFFF", size=9)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            celda.border = BORDE
        if col in COLUMNAS_SOLO_VALIDAR:
            ws.cell(2, c).comment = Comment("Formula en CONTABILIDAD: solo para validar, no pegar encima.", AUTOR)
        ws.column_dimensions[col].width = 11 if col == "A" else 12
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "B3"


def _celda_fecha(ws, i, dia, hoja, h):
    """Escribe la fecha; devuelve False si el dia no tiene fila (no hay datos)."""
    fecha = dt.date.fromisoformat(dia["fecha"]) if isinstance(dia["fecha"], str) else dia["fecha"]
    cf = ws.cell(i, 1, fecha)
    cf.number_format = "d-mm-yy"
    cf.font = Font(bold=True)
    cf.border = BORDE
    if h["fila"] is None:
        cf.fill = AMARILLO
        cf.comment = Comment("CONTABILIDAD no tiene fila para esta fecha", AUTOR)
        return False
    avisos = [a for a in dia["avisos"] if not a.startswith(("JP:", "RC:")) or a.startswith(f"{hoja}:")]
    if h.get("nueva"):
        avisos = ["Fila nueva: no estaba en CONTABILIDAD, el programa la llena", *avisos]
    if avisos:
        cf.fill = AZUL if h.get("nueva") else AMARILLO
        cf.comment = Comment("\n".join(avisos)[:2000], AUTOR)
    return True


def _celda_dato(celda, col, h, difs, pregunta):
    celda.border = BORDE
    if col in ("AB", "AC") and pregunta and pregunta["pendiente"]:
        celda.fill = AMARILLO
        celda.comment = Comment(
            f"Dia de comite: Caja 3 trae {pregunta['total']:,.2f} (alquiler + inversion en motos).\n"
            "Falta digitar el alquiler total en la app para separarlo.", AUTOR)
        return
    if col in COLUMNAS_FORMULA or col not in h["esperado"]:
        celda.fill = GRIS
        return
    celda.value = round(h["esperado"][col], 2)
    celda.number_format = FORMATO_NUM
    d = difs.get(col)
    if d:
        celda.fill = ROJO
        celda.font = Font(bold=True, color="9C0006")
        excel = "vacia" if d["excel"] is None else f"{d['excel']:,.2f}"
        celda.comment = Comment(f"En CONTABILIDAD ({col}{h['fila']}): {excel}\nDiferencia: {d['diferencia']:+,.2f}", AUTOR)
    else:
        celda.fill = AZUL if h.get("nueva") else VERDE
    if col in COLUMNAS_SOLO_VALIDAR:
        celda.font = Font(italic=True, bold=bool(d), color="9C0006" if d else "595959")


def generar_excel(dias, encabezados=None):
    """`dias` = lista de resultados de validador.validar_fecha. Devuelve bytes .xlsx."""
    wb = Workbook()
    wb.remove(wb.active)
    for hoja in HOJAS:
        ws = wb.create_sheet(hoja)
        _escribir_encabezados(ws, hoja, encabezados)
        for i, dia in enumerate(dias, start=3):
            h = dia["hojas"][hoja]
            if not _celda_fecha(ws, i, dia, hoja, h):
                continue
            difs = {d["columna"]: d for d in h["diferencias"]}
            for c in range(2, N_COLS + 1):
                _celda_dato(ws.cell(i, c), get_column_letter(c), h, difs, h.get("pregunta_alquiler"))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
